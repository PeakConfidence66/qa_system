"""
批量导入对话框 - 支持 CSV / Excel 双格式自适应导入
"""
from PyQt6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QTableWidget, QTableWidgetItem,
    QFileDialog, QMessageBox, QHeaderView,
    QProgressBar, QCheckBox, QAbstractItemView
)
from PyQt6.QtCore import Qt
import csv
import os

from models import Database, School, College, Link


class BatchImportDialog(QDialog):
    """批量导入对话框"""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("批量导入学校/学院/链接")
        self.setMinimumSize(900, 600)
        self.db = Database()
        self.csv_data = []  # 兼容性变量名：存放最终解析出的结构化字典列表
        self._csv_header = []  # 存放原始表头
        self._setup_ui()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(15)
        layout.setContentsMargins(20, 20, 20, 20)

        # 文件选择区域
        file_layout = QHBoxLayout()
        file_label = QLabel("导入文件:")
        self.file_path_label = QLabel("未选择文件")
        self.file_path_label.setStyleSheet("color: #666;")

        self.select_file_btn = QPushButton("选择数据文件")
        self.select_file_btn.setStyleSheet("""
            QPushButton {
                background-color: #1890ff;
                color: white;
                border: none;
                padding: 6px 15px;
                border-radius: 4px;
            }
            QPushButton:hover {
                background-color: #40a9ff;
            }
        """)
        self.select_file_btn.clicked.connect(self._select_file)

        file_layout.addWidget(file_label)
        file_layout.addWidget(self.file_path_label, 1)
        file_layout.addWidget(self.select_file_btn)
        layout.addLayout(file_layout)

        # 预览表格
        preview_label = QLabel("数据预览:")
        preview_label.setStyleSheet("font-weight: bold;")
        layout.addWidget(preview_label)

        self.table = QTableWidget()
        self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels([
            "源文件", "学校名称", "学院名称", "类别",
            "学院链接", "主链接", "副链接"
        ])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        layout.addWidget(self.table)

        # 导入选项
        options_layout = QHBoxLayout()

        self.skip_existing_school_cb = QCheckBox("跳过已存在的学校")
        self.skip_existing_school_cb.setChecked(True)
        options_layout.addWidget(self.skip_existing_school_cb)

        self.skip_existing_college_cb = QCheckBox("跳过已存在的学院")
        self.skip_existing_college_cb.setChecked(False)
        options_layout.addWidget(self.skip_existing_college_cb)

        options_layout.addStretch()

        self.record_count_label = QLabel("共 0 条记录")
        options_layout.addWidget(self.record_count_label)

        layout.addLayout(options_layout)

        # 进度条
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)

        # 按钮区域
        btn_layout = QHBoxLayout()
        btn_layout.addStretch()

        self.import_btn = QPushButton("开始导入")
        self.import_btn.setStyleSheet("""
            QPushButton {
                background-color: #52c41a;
                color: white;
                border: none;
                padding: 8px 25px;
                border-radius: 4px;
            }
            QPushButton:hover {
                background-color: #73d13d;
            }
            QPushButton:disabled {
                background-color: #b7eb8f;
            }
        """)
        self.import_btn.clicked.connect(self._do_import)
        self.import_btn.setEnabled(False)

        self.cancel_btn = QPushButton("取消")
        self.cancel_btn.setStyleSheet("""
            QPushButton {
                background-color: #f0f0f0;
                border: 1px solid #d9d9d9;
                padding: 8px 25px;
                border-radius: 4px;
            }
            QPushButton:hover {
                background-color: #e6e6e6;
            }
        """)
        self.cancel_btn.clicked.connect(self.reject)

        btn_layout.addWidget(self.import_btn)
        btn_layout.addWidget(self.cancel_btn)
        layout.addLayout(btn_layout)

    def _select_file(self):
        """自适应选择数据文件（支持 CSV 与 XLSX）"""
        file_path, _ = QFileDialog.getOpenFileName(
            self, "选择数据文件", "",
            "数据表格文件 (*.csv *.xlsx);;CSV文件 (*.csv);;Excel工作簿 (*.xlsx);;所有文件 (*.*)"
        )

        if file_path:
            if file_path.lower().endswith('.xlsx'):
                self._load_excel(file_path)
            else:
                self._load_csv(file_path)

    def _parse_row_data(self, header, cleaned_row):
        """【新增核心】抽象通用行解析器，自适应表头匹配（CSV 与 Excel 共同复用）"""

        def get_col_idx(name):
            try:
                return header.index(name)
            except ValueError:
                return -1

        school_idx = get_col_idx('学校名称') if get_col_idx('学校名称') >= 0 else get_col_idx('学校')
        if school_idx < 0:
            school_idx = 1

        college_idx = get_col_idx('学院名称') if get_col_idx('学院名称') >= 0 else get_col_idx('学院')
        if college_idx < 0:
            college_idx = 2

        category_idx = get_col_idx('类别')
        if category_idx < 0:
            category_idx = 3

        college_url_idx = get_col_idx('学院链接')
        if college_url_idx < 0:
            college_url_idx = 4

        main_idx = get_col_idx('主链接')
        if main_idx < 0:
            main_idx = 5

        sub_idx = get_col_idx('副链接')
        if sub_idx < 0:
            sub_idx = 6

        source_idx = get_col_idx('源文件')

        return {
            '_raw_row': cleaned_row,  # 保存原始清洗后的行数组用于表格预览
            'source_file': cleaned_row[source_idx] if 0 <= source_idx < len(cleaned_row) else '',
            'school_name': cleaned_row[school_idx] if 0 <= school_idx < len(cleaned_row) else '',
            'college_name': cleaned_row[college_idx] if 0 <= college_idx < len(cleaned_row) else '',
            'category': cleaned_row[category_idx] if 0 <= category_idx < len(cleaned_row) else '',
            'college_url': cleaned_row[college_url_idx] if 0 <= college_url_idx < len(cleaned_row) else '',
            'main_url': cleaned_row[main_idx] if 0 <= main_idx < len(cleaned_row) else '',
            'sub_url': cleaned_row[sub_idx] if 0 <= sub_idx < len(cleaned_row) else ''
        }

    def _load_excel(self, file_path):
        """【新增核心】加载 Excel 文件并进行结构化清洗"""
        try:
            # 延迟安全导入 openpyxl
            from openpyxl import load_workbook
        except ImportError:
            QMessageBox.critical(
                self, "错误",
                "导入Excel需要安装 openpyxl 依赖库。\n请在终端中执行：pip install openpyxl"
            )
            return

        try:
            self.csv_data = []

            # 以只读及计算值模式打开工作簿（效率更高）
            wb = load_workbook(file_path, read_only=True, data_only=True)
            sheet = wb.active  # 获取第一个活跃的 sheet

            # 读取所有行数据
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                QMessageBox.warning(self, "警告", "Excel文件为空，未找到任何数据")
                wb.close()
                return

            # 1. 提取并清理表头
            header = [str(cell).strip() for cell in rows[0] if cell is not None]
            self._csv_header = header

            # 2. 遍历并解析数据行
            for row in rows[1:]:
                # 过滤全空行
                if not any(row):
                    continue

                # 清洗单元格数据
                cleaned_row = [str(cell).strip().strip('`') if cell is not None else "" for cell in row]

                # 调用复用解析器
                item = self._parse_row_data(header, cleaned_row)

                # 验证学校名必填属性
                if item['school_name']:
                    self.csv_data.append(item)

            wb.close()  # 安全关闭

            # 更新UI预览
            self.file_path_label.setText(os.path.basename(file_path))
            self._update_preview_table()

            if self.csv_data:
                self.import_btn.setEnabled(True)
            else:
                QMessageBox.warning(self, "警告", "Excel文件中没有找到有效学校数据")

        except Exception as e:
            QMessageBox.critical(self, "错误", f"读取Excel文件失败: {str(e)}")

    def _load_csv(self, file_path):
        """加载CSV文件"""
        try:
            self.csv_data = []
            encodings = ['utf-8', 'utf-8-sig', 'gbk', 'gb2312']

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        reader = csv.reader(f)

                        # 跳过并解析表头
                        header = next(reader)
                        header = [h.strip() for h in header]
                        self._csv_header = header

                        for row in reader:
                            if len(row) >= 3:
                                cleaned_row = [cell.strip().strip('`') for cell in row]

                                # 调用复用解析器
                                item = self._parse_row_data(header, cleaned_row)

                                if item['school_name']:
                                    self.csv_data.append(item)
                    break  # 成功加载，退出编码循环

                except (UnicodeDecodeError, ValueError):
                    continue

            # 更新UI
            self.file_path_label.setText(os.path.basename(file_path))
            self._update_preview_table()

            if self.csv_data:
                self.import_btn.setEnabled(True)
            else:
                QMessageBox.warning(self, "警告", "CSV文件中没有找到有效数据")

        except Exception as e:
            QMessageBox.critical(self, "错误", f"加载CSV文件失败: {str(e)}")

    def _update_preview_table(self):
        """更新预览表格（动态自适应列）"""
        if not self.csv_data:
            return

        # 使用动态表头列数
        col_count = len(self._csv_header) if self._csv_header else 7
        self.table.setColumnCount(col_count)
        if self._csv_header:
            self.table.setHorizontalHeaderLabels(self._csv_header)

        self.table.setRowCount(len(self.csv_data))

        for i, item in enumerate(self.csv_data):
            # 优先显示干净、对应原始行位置的各列数据
            if '_raw_row' in item and item['_raw_row']:
                for j, val in enumerate(item['_raw_row']):
                    self.table.setItem(i, j, QTableWidgetItem(str(val)))
            else:
                # 兼容旧数据缓存格式
                self.table.setItem(i, 0, QTableWidgetItem(item.get('source_file', '')))
                self.table.setItem(i, 1, QTableWidgetItem(item.get('school_name', '')))
                self.table.setItem(i, 2, QTableWidgetItem(item.get('college_name', '')))
                self.table.setItem(i, 3, QTableWidgetItem(item.get('category', '')))
                self.table.setItem(i, 4, QTableWidgetItem(item.get('college_url', '')))
                self.table.setItem(i, 5, QTableWidgetItem(item.get('main_url', '')))
                self.table.setItem(i, 6, QTableWidgetItem(item.get('sub_url', '')))

        self.record_count_label.setText(f"共 {len(self.csv_data)} 条记录")

    def _do_import(self):
        """执行导入流程"""
        if not self.csv_data:
            return

        skip_existing_school = self.skip_existing_school_cb.isChecked()
        skip_existing_college = self.skip_existing_college_cb.isChecked()

        self.progress_bar.setVisible(True)
        self.progress_bar.setMaximum(len(self.csv_data))
        self.progress_bar.setValue(0)
        self.import_btn.setEnabled(False)

        imported_count = 0
        skipped_count = 0
        error_count = 0

        for i, item in enumerate(self.csv_data):
            try:
                # 处理学校
                school = self.db.get_school_by_name(item['school_name'])

                if school:
                    if skip_existing_school:
                        skipped_count += 1
                        self.progress_bar.setValue(i + 1)
                        continue
                    school_id = school.id
                else:
                    # 创建新学校
                    school = School(
                        name=item['school_name'],
                        category=item['category'] or '普通'
                    )
                    school_id = self.db.add_school(school)

                # 处理学院
                college = self.db.get_college_by_name(school_id, item['college_name'])

                if college:
                    if skip_existing_college:
                        skipped_count += 1
                        self.progress_bar.setValue(i + 1)
                        continue
                    college_id = college.id
                else:
                    # 创建新学院
                    college = College(
                        school_id=school_id,
                        name=item['college_name']
                    )
                    college_id = self.db.add_college(college)

                # 添加链接
                # 主链接
                if item['main_url']:
                    link = Link(
                        college_id=college_id,
                        url=item['main_url'],
                        link_type='main'
                    )
                    link.set_rules({
                        "list_selector": "",
                        "title_selector": "",
                        "link_selector": "",
                        "date_selector": "",
                        "date_format": "%Y-%m-%d"
                    })
                    self.db.add_link(link)

                # 副链接
                if item['sub_url']:
                    link = Link(
                        college_id=college_id,
                        url=item['sub_url'],
                        link_type='sub'
                    )
                    link.set_rules({
                        "list_selector": "",
                        "title_selector": "",
                        "link_selector": "",
                        "date_selector": "",
                        "date_format": "%Y-%m-%d"
                    })
                    self.db.add_link(link)

                imported_count += 1

            except Exception as e:
                error_count += 1
                print(f"导入失败: {item['school_name']} - {str(e)}")

            self.progress_bar.setValue(i + 1)

        self.progress_bar.setVisible(False)

        # 显示结果
        msg = f"导入完成！\n\n成功导入: {imported_count} 条\n跳过: {skipped_count} 条\n失败: {error_count} 条"

        if error_count > 0:
            QMessageBox.warning(self, "导入完成", msg)
        else:
            QMessageBox.information(self, "导入成功", msg)

        if imported_count > 0:
            self.accept()

    def closeEvent(self, event):
        """关闭事件"""
        self.db.close()
        event.accept()